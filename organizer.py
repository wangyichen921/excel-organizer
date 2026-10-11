"""合并表头相同的 Excel 文件；不修改输入文件。"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from tempfile import NamedTemporaryFile
import os

from openpyxl import Workbook, load_workbook


@dataclass(frozen=True)
class MergeResult:
    input_rows: int
    output_rows: int
    duplicate_rows: int
    output_path: Path


def read_header(path: str | Path) -> tuple[str, ...]:
    """读取第一个工作表的表头，并拒绝含空列名或重复列名的表格。"""
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book.worksheets[0]
        cells = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header = tuple(cells)
        if not header or any(not isinstance(cell, str) or not cell.strip() for cell in header):
            raise ValueError(f"表头为空或包含空列名：{path}")
        if len(set(header)) != len(header):
            raise ValueError(f"表头包含重复列名：{path}")
        return header
    finally:
        book.close()


def _row_key(row: tuple) -> tuple:
    # 把类型纳入去重键，避免 True 与 1 被 Python 当成同一个值。
    return tuple((type(value), value) for value in row)


def _sort_key(value: object) -> tuple:
    # 不同类型分组排序；空值在升序和降序时都放在末尾。
    if isinstance(value, bool):
        return (3, str(value))
    if isinstance(value, (int, float)):
        return (0, value)
    if isinstance(value, (datetime, date, time)):
        return (1, value.isoformat())
    return (2, str(value).casefold())


def merge_workbooks(
    input_paths: list[str | Path],
    output_path: str | Path,
    sort_column: str,
    descending: bool = False,
) -> MergeResult:
    """按文件选择顺序合并首个工作表，整行去重后按指定列排序。"""
    if not input_paths:
        raise ValueError("请至少选择一个 .xlsx 文件。")

    paths = [Path(path).resolve() for path in input_paths]
    output = Path(output_path).resolve()
    if any(path.suffix.lower() != ".xlsx" or not path.is_file() for path in paths):
        raise ValueError("输入文件必须是存在的 .xlsx 文件。")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("输出文件必须使用 .xlsx 扩展名。")
    if output in paths:
        raise ValueError("输出路径不能与输入文件相同，以免覆盖原文件。")

    header = read_header(paths[0])
    if sort_column not in header:
        raise ValueError(f"排序列不存在：{sort_column}")

    rows: list[tuple] = []
    seen: set[tuple] = set()
    input_count = 0
    for path in paths:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            current_header = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
            if current_header != header:
                raise ValueError(f"表头不一致：{path.name}；要求列名和顺序完全相同。")
            for values in sheet.iter_rows(min_row=2, values_only=True):
                row = tuple(values[: len(header)])
                if all(value is None for value in row):
                    continue
                input_count += 1
                key = _row_key(row)
                if key not in seen:
                    seen.add(key)
                    rows.append(row)
        finally:
            book.close()

    column_index = header.index(sort_column)
    present = [row for row in rows if row[column_index] is not None]
    missing = [row for row in rows if row[column_index] is None]
    present.sort(key=lambda row: _sort_key(row[column_index]), reverse=descending)
    rows = present + missing

    output.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "整理结果"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    temporary: Path | None = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        os.replace(temporary, output)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)

    return MergeResult(input_count, len(rows), input_count - len(rows), output)
