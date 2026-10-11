from pathlib import Path
from tempfile import TemporaryDirectory
from datetime import datetime
import unittest

from openpyxl import Workbook, load_workbook

from organizer import merge_workbooks


def write_book(path: Path, header: tuple, rows: list[tuple]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()


class OrganizerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = TemporaryDirectory(dir=Path(__file__).resolve().parent)
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.first = self.folder / "一.xlsx"
        self.second = self.folder / "二.xlsx"
        self.output = self.folder / "结果.xlsx"

    def output_rows(self) -> list[tuple]:
        book = load_workbook(self.output, read_only=True, data_only=True)
        try:
            return list(book.active.values)
        finally:
            book.close()

    def test_merge_deduplicate_sort_and_preserve_sources(self) -> None:
        write_book(self.first, ("编号", "分数"), [("甲", 2), ("乙", 1)])
        write_book(self.second, ("编号", "分数"), [("甲", 2), ("丙", 3)])
        original_first = self.first.read_bytes()
        original_second = self.second.read_bytes()

        result = merge_workbooks([self.first, self.second], self.output, "分数")

        self.assertEqual((result.input_rows, result.duplicate_rows, result.output_rows), (4, 1, 3))
        self.assertEqual(self.output_rows(), [("编号", "分数"), ("乙", 1), ("甲", 2), ("丙", 3)])
        self.assertEqual(self.first.read_bytes(), original_first)
        self.assertEqual(self.second.read_bytes(), original_second)

    def test_descending_keeps_blank_values_last(self) -> None:
        write_book(self.first, ("编号", "分数"), [("甲", None), ("乙", 1), ("丙", 3)])
        merge_workbooks([self.first], self.output, "分数", descending=True)
        self.assertEqual(self.output_rows()[1:], [("丙", 3), ("乙", 1), ("甲", None)])

    def test_date_column_sort(self) -> None:
        write_book(self.first, ("编号", "日期"), [
            ("晚", datetime(2026, 1, 3)),
            ("早", datetime(2026, 1, 1)),
        ])
        merge_workbooks([self.first], self.output, "日期")
        self.assertEqual([row[0] for row in self.output_rows()[1:]], ["早", "晚"])

    def test_header_mismatch_does_not_create_output(self) -> None:
        write_book(self.first, ("编号", "分数"), [("甲", 1)])
        write_book(self.second, ("分数", "编号"), [(1, "甲")])
        with self.assertRaisesRegex(ValueError, "表头不一致"):
            merge_workbooks([self.first, self.second], self.output, "分数")
        self.assertFalse(self.output.exists())

    def test_refuses_to_overwrite_source(self) -> None:
        write_book(self.first, ("编号",), [("甲",)])
        with self.assertRaisesRegex(ValueError, "输出路径不能与输入文件相同"):
            merge_workbooks([self.first], self.first, "编号")

    def test_missing_sort_column(self) -> None:
        write_book(self.first, ("编号",), [("甲",)])
        with self.assertRaisesRegex(ValueError, "排序列不存在"):
            merge_workbooks([self.first], self.output, "金额")

    def test_duplicate_header_is_rejected(self) -> None:
        write_book(self.first, ("编号", "编号"), [("甲", "乙")])
        with self.assertRaisesRegex(ValueError, "重复列名"):
            merge_workbooks([self.first], self.output, "编号")


if __name__ == "__main__":
    unittest.main()
